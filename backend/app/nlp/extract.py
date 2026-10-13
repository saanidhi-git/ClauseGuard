# Read text from PDF, DOCX or TXT

import io
import re
from collections import Counter
from dataclasses import dataclass, field
from typing import List, Tuple

from ..config import MAX_CHARS, MAX_PAGES


# friendly error shown to the user
class ExtractError(ValueError):
    pass


@dataclass
class Extracted:
    text: str
    page_spans: List[Tuple[int, int, int]] = field(default_factory=list)
    page_count: int = 1

    # page number for a position in the text
    def page_at(self, pos: int) -> int:
        for start, end, page in self.page_spans:
            if start <= pos < end:
                return page
        return self.page_spans[-1][2] if self.page_spans else 1


PAGE_NUM_RE = re.compile(r"^\s*(page\s*)?\d{1,3}(\s*(of|/)\s*\d{1,3})?\s*$", re.I)


def _norm_line(line: str) -> str:
    return re.sub(r"\d+", "#", line.strip().lower())


# remove headers, footers and page numbers
def _clean_pages(pages: List[str]) -> List[str]:
    split = [p.splitlines() for p in pages]
    repeated = set()
    if len(pages) >= 3:
        counts: Counter = Counter()
        for lines in split:
            edge = [l for l in lines if l.strip()]
            edge = edge[:2] + edge[-2:]
            counts.update({_norm_line(l) for l in edge})
        repeated = {l for l, c in counts.items() if c >= max(2, len(pages) * 0.6) and l}

    cleaned = []
    for lines in split:
        keep = [l for l in lines if not PAGE_NUM_RE.match(l) and _norm_line(l) not in repeated]
        text = "\n".join(keep)
        text = re.sub(r"(\w)-\n(\w)", r"\1\2", text)  # join hyphenated words
        cleaned.append(text)
    return cleaned


def _finalize(text: str) -> str:
    text = text.replace("\r", "").replace(" ", " ").replace("\t", " ")
    text = re.sub(r"[ ]{2,}", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _join_pages(pages: List[str]) -> Extracted:
    parts, spans, pos = [], [], 0
    for i, page in enumerate(pages, start=1):
        page = _finalize(page)
        if not page:
            continue
        if parts:
            parts.append("\n\n")
            pos += 2
        spans.append((pos, pos + len(page), i))
        parts.append(page)
        pos += len(page)
    return Extracted(text="".join(parts), page_spans=spans, page_count=max(1, len(pages)))


def from_pdf(data: bytes) -> Extracted:
    import pymupdf as fitz

    try:
        doc = fitz.open(stream=data, filetype="pdf")
    except Exception as exc:
        raise ExtractError("This PDF could not be opened. It may be damaged or password-protected.") from exc
    if doc.page_count > MAX_PAGES:
        raise ExtractError(f"This document has {doc.page_count} pages. The limit is {MAX_PAGES} pages.")
    pages = [page.get_text("text") for page in doc]
    doc.close()
    if sum(len(p.strip()) for p in pages) < 30:
        raise ExtractError("This looks like a scanned PDF; OCR isn't supported yet.")
    return _join_pages(_clean_pages(pages))


def from_docx(data: bytes) -> Extracted:
    import docx

    try:
        d = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise ExtractError("This DOCX file could not be read.") from exc
    lines = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            lines.append(" | ".join(c.text.strip() for c in row.cells))
    return _join_pages(["\n".join(lines)])


def from_text(text: str) -> Extracted:
    return _join_pages([text or ""])


# read any supported file
def extract(data: bytes, filename: str) -> Extracted:
    name = (filename or "").lower()
    if name.endswith(".pdf"):
        result = from_pdf(data)
    elif name.endswith(".docx"):
        result = from_docx(data)
    elif name.endswith(".txt") or name.endswith(".md"):
        result = from_text(data.decode("utf-8", errors="ignore"))
    else:
        raise ExtractError("Unsupported file type. Please upload a PDF, DOCX or TXT file.")
    return validate(result)


# reject empty or very long documents
def validate(result: Extracted) -> Extracted:
    if len(result.text.split()) < 5:
        raise ExtractError("The document is empty or too short to analyse. Please check the file or paste the full text.")
    if len(result.text) > MAX_CHARS:
        raise ExtractError("The document is too long. Please upload something under about 40 pages.")
    return result
